"""文档读取离线测试：不发网络请求，也不启动 Word。"""

from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from document_loader import (
    DocumentLoadError,
    combine_project_text,
    load_document,
    load_uploaded_document,
)


def pdf_bytes_with_text(text: bytes = b"Hello PDF") -> bytes:
    """构造只有一页且包含文本层的 PDF，不依赖额外测试工具。"""
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (" + text + b") Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(output)
    return output.getvalue()


class DocumentLoaderTests(unittest.TestCase):
    def test_markdown_path_and_utf8_upload(self):
        with TemporaryDirectory() as directory:
            sample = Path(directory) / "项目.MD"
            sample.write_text("# 项目\n使用 Redis。", encoding="utf-8")
            document = load_document(sample)
        self.assertEqual(document.source, "项目.MD")
        self.assertEqual(document.file_type, "md")
        self.assertIn("使用 Redis", document.segments[0].text)

        uploaded = load_uploaded_document("dir\\笔记.md", b"\xef\xbb\xbf" + "中文".encode("utf-8"))
        self.assertEqual(uploaded.source, "笔记.md")
        self.assertEqual(uploaded.segments[0].text, "中文")

    def test_pdf_with_text_and_source_page(self):
        document = load_uploaded_document("resume.pdf", pdf_bytes_with_text())
        self.assertEqual(document.segments[0].text, "Hello PDF")
        self.assertEqual(document.segments[0].page, 1)
        self.assertIn("[来源：resume.pdf，第 1 页]", combine_project_text([document]))

    def test_docx_paragraphs_header_and_tables(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "简历页眉"
        document.add_paragraph("项目经验")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "技术"
        table.cell(0, 1).text = "Redis"
        output = BytesIO()
        document.save(output)

        loaded = load_uploaded_document("简历.docx", output.getvalue())
        self.assertEqual(loaded.file_type, "docx")
        self.assertIn("简历页眉\n项目经验\n技术 | Redis", loaded.segments[0].text)

    def test_legacy_doc_uses_isolated_word_and_closes_it(self):
        fake_document = SimpleNamespace(Content=SimpleNamespace(Text="旧版简历内容"), Close=Mock())
        fake_word = SimpleNamespace(Documents=SimpleNamespace(Open=Mock(return_value=fake_document)), Quit=Mock())
        fake_client = ModuleType("win32com.client")
        fake_client.DispatchEx = Mock(return_value=fake_word)
        fake_package = ModuleType("win32com")
        fake_package.client = fake_client
        signature = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

        with (
            patch("document_loader.sys.platform", "win32"),
            patch.dict(sys.modules, {"win32com": fake_package, "win32com.client": fake_client, "pythoncom": Mock()}),
        ):
            loaded = load_uploaded_document("旧版.doc", signature + b"stub")

        self.assertEqual(loaded.segments[0].text, "旧版简历内容")
        self.assertEqual(loaded.file_type, "doc")
        fake_client.DispatchEx.assert_called_once_with("Word.Application")
        self.assertEqual(fake_word.AutomationSecurity, 3)
        self.assertTrue(fake_word.Documents.Open.call_args.kwargs["ReadOnly"])
        fake_document.Close.assert_called_once_with(SaveChanges=0)
        fake_word.Quit.assert_called_once_with()

    def test_empty_unsupported_invalid_and_unreadable_files(self):
        for filename, data in (
            ("a.md", b"   "), ("a.md", b"\xff"), ("a.pdf", b"not a pdf"),
            ("a.doc", b"not a doc"), ("a.docx", b"not a zip"),
            ("a.zip", b"any"), ("a.md", b""),
        ):
            with self.subTest(filename=filename, data=data), self.assertRaises(DocumentLoadError):
                load_uploaded_document(filename, data)

        with self.assertRaises(DocumentLoadError):
            load_document("missing-material.md")

    def test_pdf_without_text_is_rejected(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(DocumentLoadError, "未提取到文字"):
            load_uploaded_document("扫描简历.pdf", output.getvalue())

    def test_combining_documents_keeps_sources(self):
        resume = load_uploaded_document("简历.pdf", pdf_bytes_with_text())
        project = load_uploaded_document("项目.md", "使用 Redis。".encode("utf-8"))
        combined = combine_project_text([resume, project])
        self.assertIn("[来源：简历.pdf，第 1 页]", combined)
        self.assertIn("[来源：项目.md]", combined)
        self.assertLess(combined.index("Hello PDF"), combined.index("使用 Redis"))
        with self.assertRaises(DocumentLoadError):
            combine_project_text([])


if __name__ == "__main__":
    unittest.main()
