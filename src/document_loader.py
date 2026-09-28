"""将本地或上传的 PDF、Markdown、DOC、DOCX 转成带来源的文本。"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PureWindowsPath
import sys
from tempfile import TemporaryDirectory
from typing import Sequence
from zipfile import BadZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


SUPPORTED_SUFFIXES = {".md", ".pdf", ".doc", ".docx"}
MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_TOTAL_CHARS = 120_000


class DocumentLoadError(ValueError):
    """文件无效、解析失败或无法提取有效文字。"""


@dataclass(frozen=True)
class TextSegment:
    text: str
    page: int | None = None


@dataclass(frozen=True)
class LoadedDocument:
    source: str
    file_type: str
    segments: tuple[TextSegment, ...]


def load_document(path: str | Path) -> LoadedDocument:
    """读取本地材料；文档不写入项目目录。"""
    file_path = Path(path)
    if not file_path.is_file():
        raise DocumentLoadError(f"文件不存在或不是普通文件：{file_path}")
    if file_path.stat().st_size > MAX_FILE_BYTES:
        raise DocumentLoadError(f"文件超过 {MAX_FILE_BYTES // (1024 * 1024)} MB 上限：{file_path.name}")
    return load_uploaded_document(file_path.name, file_path.read_bytes())


def load_uploaded_document(filename: str, data: bytes) -> LoadedDocument:
    """解析上传的文件字节；可直接接 Streamlit UploadedFile.getvalue()。"""
    if not isinstance(filename, str) or not filename.strip():
        raise DocumentLoadError("文件名不能为空")

    source = PureWindowsPath(filename).name
    suffix = Path(source).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentLoadError("仅支持 .md、.pdf、.doc、.docx 文件")
    if not isinstance(data, bytes) or not data:
        raise DocumentLoadError(f"{source} 为空或不是二进制文件内容")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentLoadError(f"{source} 超过 {MAX_FILE_BYTES // (1024 * 1024)} MB 上限")

    if suffix == ".md":
        try:
            segments = (TextSegment(data.decode("utf-8-sig").strip()),)
        except UnicodeDecodeError as exc:
            raise DocumentLoadError(f"{source} 不是 UTF-8 编码的 Markdown") from exc
    elif suffix == ".pdf":
        segments = _extract_pdf(data, source)
    elif suffix == ".docx":
        segments = (TextSegment(_extract_docx(data, source)),)
    else:
        segments = (TextSegment(_extract_doc(data, source)),)

    if not any(segment.text.strip() for segment in segments):
        raise DocumentLoadError(f"{source} 未提取到文字；扫描件 PDF 暂不支持 OCR")
    if sum(len(segment.text) for segment in segments) > MAX_TOTAL_CHARS:
        raise DocumentLoadError(f"{source} 提取出的文字超过 {MAX_TOTAL_CHARS} 字符上限")

    return LoadedDocument(source, suffix.removeprefix("."), segments)


def _extract_pdf(data: bytes, source: str) -> tuple[TextSegment, ...]:
    if not data.startswith(b"%PDF-"):
        raise DocumentLoadError(f"{source} 不是有效的 PDF 文件")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentLoadError(f"{source} 已加密，请先移除密码")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentLoadError(f"{source} 超过 {MAX_PDF_PAGES} 页上限")
        return tuple(
            TextSegment(text=text, page=page_number)
            for page_number, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        )
    except (PdfReadError, OSError, ValueError) as exc:
        if isinstance(exc, DocumentLoadError):
            raise
        raise DocumentLoadError(f"无法解析 PDF：{source}") from exc


def _extract_docx(data: bytes, source: str) -> str:
    try:
        document = Document(BytesIO(data))
        parts = []
        # 简历的姓名或联系信息有时位于页眉。
        for section in document.sections:
            parts.extend(paragraph.text.strip() for paragraph in section.header.paragraphs if paragraph.text.strip())
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph) and block.text.strip():
                parts.append(block.text.strip())
            elif isinstance(block, Table):
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        return "\n".join(parts)
    except (BadZipFile, PackageNotFoundError, OSError, ValueError) as exc:
        raise DocumentLoadError(f"无法解析 DOCX：{source}") from exc


def _extract_doc(data: bytes, source: str) -> str:
    """使用独立的 Word COM 进程读取旧版二进制 DOC，不修改用户文档。"""
    if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise DocumentLoadError(f"{source} 不是有效的旧版 DOC 文件")
    if sys.platform != "win32":
        raise DocumentLoadError(".doc 解析需要 Windows 与 Microsoft Word；可先另存为 .docx")
    try:
        from win32com.client import DispatchEx
    except ImportError as exc:
        raise DocumentLoadError(".doc 解析需要 pywin32 和 Microsoft Word") from exc

    try:
        with TemporaryDirectory(prefix="mock-interview-doc-") as directory:
            temporary_doc = Path(directory) / "source.doc"
            temporary_doc.write_bytes(data)
            word = None
            document = None
            try:
                word = DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                word.AutomationSecurity = 3  # 禁用文档中的宏
                document = word.Documents.Open(
                    str(temporary_doc), ReadOnly=True, AddToRecentFiles=False,
                    ConfirmConversions=False, UpdateLinks=0,
                )
                return document.Content.Text.strip()
            finally:
                # 在 TemporaryDirectory 清理前释放 Word 对文件的占用。
                try:
                    if document is not None:
                        document.Close(SaveChanges=0)
                finally:
                    if word is not None:
                        word.Quit()
    except Exception as exc:
        raise DocumentLoadError(f"无法用 Microsoft Word 读取 DOC：{source}") from exc


def combine_project_text(documents: Sequence[LoadedDocument]) -> str:
    """将多份材料组合成当前 Claim Extractor 所需的 project_text。"""
    if not documents:
        raise DocumentLoadError("至少提供一份材料")
    parts = []
    for document in documents:
        for segment in document.segments:
            location = f"，第 {segment.page} 页" if segment.page is not None else ""
            parts.append(f"[来源：{document.source}{location}]\n{segment.text}")
    project_text = "\n\n".join(parts)
    if len(project_text) > MAX_TOTAL_CHARS:
        raise DocumentLoadError(f"合并后的材料超过 {MAX_TOTAL_CHARS} 字符上限")
    return project_text
